#!/usr/bin/env python3
"""Create a formatting-light DOCX so Word translates text, not fragmented runs."""
import json
import sys
from pathlib import Path

from docx import Document


def main() -> int:
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    data = json.loads(source.read_text(encoding="utf-8"))
    doc = Document()
    # Remove the starter paragraph; every source paragraph is recreated once.
    if doc.paragraphs:
        body = doc._element.body
        body.remove(doc.paragraphs[0]._element)
    for item in data["paragraphs"]:
        raw = item["style"].lower()
        if raw == "title": style = "Title"
        elif raw == "heading 1": style = "Heading 1"
        elif raw == "heading 2": style = "Heading 2"
        elif raw == "heading 3": style = "Heading 3"
        elif raw == "quote": style = "Quote"
        elif "list" in raw or "bullet" in raw: style = "List Bullet"
        else: style = "Normal"
        p = doc.add_paragraph(style=style)
        p.add_run(item["text"])
    doc.save(target)
    print(target)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
