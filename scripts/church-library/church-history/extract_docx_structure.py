from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document


def extract(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.replace("\u00a0", " ").strip()
        if not text:
            continue
        paragraphs.append(
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style else None,
                "text": text,
            }
        )
    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.replace("\u00a0", " ").strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
    raw = path.read_bytes()
    return {
        "source": str(path),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "size_bytes": len(raw),
        "paragraphs": paragraphs,
        "paragraph_count": len(paragraphs),
        "tables": tables,
        "table_count": len(tables),
        "has_comments": "word/comments.xml" in names,
        "has_footnotes": "word/footnotes.xml" in names,
        "has_endnotes": "word/endnotes.xml" in names,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for path in args.inputs:
        payload = extract(path)
        out = args.output_dir / f"{path.stem}.structure.json"
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({k: payload[k] for k in ("source", "sha256", "size_bytes", "paragraph_count", "table_count", "has_comments", "has_footnotes", "has_endnotes")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
